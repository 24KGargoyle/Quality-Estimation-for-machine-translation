"""PowerPoint (.pptx) parsing via python-pptx — per-slide text, tables, and
speaker notes, preserving `slide_number` for citations."""
from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from meeting_intel.ingestion.documents import NormalizedDocument, ParseResult, UnsupportedFileError
from meeting_intel.ingestion.parsers.base import DocumentParser
from meeting_intel.ingestion.parsers.textutil import split_text


def _slide_text(slide) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
        if shape.has_table:
            table = shape.table
            headers = [c.text.strip() for c in table.rows[0].cells]
            for row in list(table.rows)[1:]:
                cells = [c.text.strip() for c in row.cells]
                line = "; ".join(f"{h}: {v}" for h, v in zip(headers, cells) if v)
                if line:
                    parts.append(line)
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            parts.append(f"Speaker notes: {notes}")
    return "\n\n".join(parts)


class PowerPointParser(DocumentParser):
    document_type = "presentation"

    def parse(
        self, *, content: bytes, filename: str, relative_path: str, tenant_id: str,
        meeting_id: str | None, title: str, document_id_prefix: str,
    ) -> ParseResult:
        try:
            prs = Presentation(io.BytesIO(content))
        except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
            raise UnsupportedFileError(f"Could not open as a .pptx file: {exc}") from exc

        documents: list[NormalizedDocument] = []
        for slide_number, slide in enumerate(prs.slides, start=1):
            text = _slide_text(slide)
            if not text.strip():
                continue
            for piece in split_text(text):
                documents.append(
                    NormalizedDocument(
                        document_id=f"{document_id_prefix}:{len(documents)}",
                        tenant_id=tenant_id,
                        meeting_id=meeting_id,
                        title=title,
                        source_file=filename,
                        relative_path=relative_path,
                        file_type="pptx",
                        document_type=self.document_type,
                        content=piece,
                        slide_number=slide_number,
                    )
                )

        warnings = [] if documents else ["No extractable text found on any slide."]
        return ParseResult(documents=documents, warnings=warnings)
