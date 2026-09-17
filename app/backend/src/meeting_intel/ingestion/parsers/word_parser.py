"""Word (.docx) parsing — paragraphs, headings (used as `section` metadata),
and tables (rendered as `Header: value` records, never flattened into a
single run-on string) are extracted in document order via python-docx.

Legacy `.doc` (the pre-2007 OLE binary format) is not something python-docx
can read — it is reported as an unsupported file (never silently ignored,
never claimed to have been parsed) rather than crashing the batch. A
sandboxed LibreOffice/`soffice --headless --convert-to docx` conversion step
is the documented production path for `.doc` (see docs/HISTORICAL_IMPORT.md)
— not implemented here since no such sandboxed converter is available in
this environment to safely exercise and verify.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from meeting_intel.ingestion.documents import NormalizedDocument, ParseResult, UnsupportedFileError
from meeting_intel.ingestion.parsers.base import DocumentParser
from meeting_intel.ingestion.parsers.textutil import split_text

_HEADING_STYLES = {f"Heading {n}" for n in range(1, 10)} | {"Title"}


def _table_to_text(table) -> str:
    rows = table.rows
    if not rows:
        return ""
    headers = [c.text.strip() for c in rows[0].cells]
    row_blocks = []
    for row in rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        lines = [f"{h}: {v}" for h, v in zip(headers, cells) if v]
        if lines:
            row_blocks.append("\n".join(lines))
    return "\n\n".join(row_blocks)


class WordParser(DocumentParser):
    document_type = "supporting_document"

    def parse(
        self, *, content: bytes, filename: str, relative_path: str, tenant_id: str,
        meeting_id: str | None, title: str, document_id_prefix: str,
    ) -> ParseResult:
        if filename.lower().endswith(".doc") and not filename.lower().endswith(".docx"):
            raise UnsupportedFileError(
                "Legacy .doc format is not supported directly — convert to .docx and re-upload."
            )
        try:
            doc = Document(io.BytesIO(content))
        except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
            raise UnsupportedFileError(f"Could not open as a .docx file: {exc}") from exc

        documents: list[NormalizedDocument] = []
        section = "Document"
        buffer_parts: list[str] = []

        def flush() -> None:
            nonlocal buffer_parts
            text = "\n\n".join(buffer_parts).strip()
            buffer_parts = []
            for piece in split_text(text):
                documents.append(
                    NormalizedDocument(
                        document_id=f"{document_id_prefix}:{len(documents)}",
                        tenant_id=tenant_id,
                        meeting_id=meeting_id,
                        title=title,
                        source_file=filename,
                        relative_path=relative_path,
                        file_type="docx",
                        document_type=self.document_type,
                        content=piece,
                        section=section,
                    )
                )

        body = doc.element.body
        table_index = 0
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = next((p for p in doc.paragraphs if p._p is child), None)
                if para is None or not para.text.strip():
                    continue
                if para.style and para.style.name in _HEADING_STYLES:
                    flush()
                    section = para.text.strip()
                else:
                    buffer_parts.append(para.text.strip())
            elif tag == "tbl":
                flush()
                if table_index < len(doc.tables):
                    table_text = _table_to_text(doc.tables[table_index])
                    table_index += 1
                    if table_text:
                        for piece in split_text(table_text):
                            documents.append(
                                NormalizedDocument(
                                    document_id=f"{document_id_prefix}:{len(documents)}",
                                    tenant_id=tenant_id,
                                    meeting_id=meeting_id,
                                    title=title,
                                    source_file=filename,
                                    relative_path=relative_path,
                                    file_type="docx",
                                    document_type=self.document_type,
                                    content=piece,
                                    section=section,
                                    metadata={"is_table": True},
                                )
                            )
        flush()

        warnings = [] if documents else ["No extractable text found in this document."]
        return ParseResult(documents=documents, warnings=warnings)
