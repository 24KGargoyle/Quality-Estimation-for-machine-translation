"""Excel (.xlsx) parsing via openpyxl in read-only, data-only mode — never
loads or executes VBA/macros (openpyxl has no macro-execution capability at
all; a `.xlsm`'s VBA project is an opaque binary blob it never touches).

Legacy `.xls` (the pre-2007 BIFF binary format) is not readable by openpyxl
and is reported as unsupported (convert to `.xlsx`) rather than crashing.

Large sheets are chunked into row batches — never one giant document per
workbook — so retrieval can return the specific rows relevant to a question
rather than the whole spreadsheet as one blob.
"""
from __future__ import annotations

import io
import zipfile

import openpyxl

from meeting_intel.ingestion.documents import NormalizedDocument, ParseResult, UnsupportedFileError
from meeting_intel.ingestion.parsers.base import DocumentParser

ROWS_PER_CHUNK = 25


def _render_batch(workbook_name: str, sheet_name: str, headers: list[str], rows: list[list[str]]) -> str:
    lines = [f"Workbook: {workbook_name}", f"Sheet: {sheet_name}", ""]
    lines.append(" | ".join(headers))
    for row in rows:
        lines.append(" | ".join(row))
    return "\n".join(lines)


class ExcelParser(DocumentParser):
    document_type = "spreadsheet"

    def parse(
        self, *, content: bytes, filename: str, relative_path: str, tenant_id: str,
        meeting_id: str | None, title: str, document_id_prefix: str,
    ) -> ParseResult:
        if filename.lower().endswith(".xls") and not filename.lower().endswith(".xlsx"):
            raise UnsupportedFileError(
                "Legacy .xls format is not supported directly — convert to .xlsx and re-upload."
            )
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise UnsupportedFileError(f"Could not open as an .xlsx workbook: {exc}") from exc

        documents: list[NormalizedDocument] = []
        for sheet in wb.worksheets:
            rows_iter = sheet.iter_rows(values_only=True)
            headers: list[str] | None = None
            batch: list[list[str]] = []
            for raw_row in rows_iter:
                cells = ["" if v is None else str(v) for v in raw_row]
                if not any(c.strip() for c in cells):
                    continue
                if headers is None:
                    headers = cells
                    continue
                batch.append(cells)
                if len(batch) >= ROWS_PER_CHUNK:
                    documents.append(
                        NormalizedDocument(
                            document_id=f"{document_id_prefix}:{sheet.title}:{len(documents)}",
                            tenant_id=tenant_id,
                            meeting_id=meeting_id,
                            title=title,
                            source_file=filename,
                            relative_path=relative_path,
                            file_type="xlsx",
                            document_type=self.document_type,
                            content=_render_batch(filename, sheet.title, headers, batch),
                            sheet_name=sheet.title,
                        )
                    )
                    batch = []
            if headers is not None and batch:
                documents.append(
                    NormalizedDocument(
                        document_id=f"{document_id_prefix}:{sheet.title}:{len(documents)}",
                        tenant_id=tenant_id,
                        meeting_id=meeting_id,
                        title=title,
                        source_file=filename,
                        relative_path=relative_path,
                        file_type="xlsx",
                        document_type=self.document_type,
                        content=_render_batch(filename, sheet.title, headers, batch),
                        sheet_name=sheet.title,
                    )
                )
        wb.close()

        warnings = [] if documents else ["No data rows found in any sheet."]
        return ParseResult(documents=documents, warnings=warnings)
