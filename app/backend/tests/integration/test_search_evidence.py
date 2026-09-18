"""Integration tests for evidence-first search/Q&A (§1 of the Search
Intelligence upgrade), exercised through the real HTTP API. §18 test cases:
generic/person/topic/document/no-result/cross-document questions.
"""
import io

import pytest
from docx import Document

from tests.conftest import SAMPLE_VTT, dev_login

pytestmark = pytest.mark.asyncio


def _sample_docx(heading: str, paragraph: str) -> bytes:
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _load_meeting(client, headers, *, meeting_id="m1", title="Standup") -> str:
    resp = await client.post(
        "/api/meetings/load", headers=headers, json={"meeting_id": meeting_id, "title": title, "transcript_vtt": SAMPLE_VTT}
    )
    return resp.json()["id"]


async def test_generic_question_returns_evidence_and_sources(client):
    token = await dev_login(client, email="a1@acme.com", display_name="A1", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    meeting_id = await _load_meeting(client, headers)

    resp = await client.post("/api/chat", headers=headers, json={"meeting_id": meeting_id, "message": "What was discussed?"})
    data = resp.json()
    assert resp.status_code == 200
    assert data["intelligence"] is not None


async def test_person_question_filters_to_that_speaker(client):
    token = await dev_login(client, email="a2@acme.com", display_name="A2", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    meeting_id = await _load_meeting(client, headers)

    resp = await client.post(
        "/api/chat", headers=headers, json={"meeting_id": meeting_id, "message": "What did Chetan say about CrewAI?"}
    )
    data = resp.json()
    assert data["speaker_filter"] == "Chetan Kumar"


async def test_document_source_is_never_labeled_as_a_transcript_speaker(client):
    """A document-sourced chunk's `document_type` must never be
    "transcript" — this is what lets the prompt/UI correctly say "the
    document identifies X" rather than "X said" (§1.5)."""
    token = await dev_login(client, email="a3@acme.com", display_name="A3", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [("files", ("Roles/Roles.docx", _sample_docx("Roles", "Contact: Alice. Role: Assignment Lead."), "application/octet-stream"))]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    job_id = create.json()["id"]

    import asyncio

    for _ in range(50):
        status_resp = await client.get(f"/api/historical-imports/{job_id}", headers=headers)
        if status_resp.json()["status"] not in ("queued", "processing"):
            break
        await asyncio.sleep(0.1)

    meetings = await client.get("/api/meetings", headers=headers)
    meeting_id = meetings.json()[0]["id"]

    resp = await client.get(f"/api/meetings/{meeting_id}/search?q=Alice Assignment Lead", headers=headers)
    results = resp.json()
    assert len(results) > 0
    assert all(r["document_type"] != "transcript" for r in results)
    assert all(r["speaker"] is None for r in results)


async def test_topic_question_returns_related_topics_in_sidebar(client):
    token = await dev_login(client, email="a4@acme.com", display_name="A4", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    meeting_id = await _load_meeting(client, headers, meeting_id="m4")

    resp = await client.get(f"/api/meetings/{meeting_id}/intelligence?q=What was decided about CrewAI?", headers=headers)
    assert resp.status_code == 200
    assert isinstance(resp.json()["related_topics"], list)


async def test_no_result_question_returns_empty_evidence_not_fabricated(client):
    token = await dev_login(client, email="a5@acme.com", display_name="A5", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    meeting_id = await _load_meeting(client, headers, meeting_id="m5")

    resp = await client.post(
        "/api/chat", headers=headers,
        json={"meeting_id": meeting_id, "message": "What did we decide about the quarterly budget for interstellar travel?"},
    )
    data = resp.json()
    assert data["evidence_sufficient"] is False
    assert data["sources"] == []


async def test_cross_document_search_returns_both_transcript_and_document_chunks(client):
    token = await dev_login(client, email="a6@acme.com", display_name="A6", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [
        ("files", ("Proj/Meeting.vtt", SAMPLE_VTT.encode(), "text/vtt")),
        (
            "files",
            ("Proj/Action_Items.docx", _sample_docx("Action Items", "Owner: Chetan. Action: Confirm CrewAI requirement."), "application/octet-stream"),
        ),
    ]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    job_id = create.json()["id"]

    import asyncio

    for _ in range(50):
        status_resp = await client.get(f"/api/historical-imports/{job_id}", headers=headers)
        if status_resp.json()["status"] not in ("queued", "processing"):
            break
        await asyncio.sleep(0.1)

    meetings = await client.get("/api/meetings", headers=headers)
    meeting_id = meetings.json()[0]["id"]

    resp = await client.get(f"/api/meetings/{meeting_id}/search?q=CrewAI requirement", headers=headers)
    results = resp.json()
    file_types = {r["file_type"] for r in results}
    assert "vtt" in file_types
    assert "docx" in file_types
