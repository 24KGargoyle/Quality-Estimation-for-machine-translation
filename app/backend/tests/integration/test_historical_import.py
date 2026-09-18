"""Integration tests for the Historical Meeting Data Import feature (§40 of
the upgrade spec), exercised through the real HTTP API end to end: upload ->
background job -> progress polling -> results -> import history -> the
resulting meeting/documents/search-index state.

Covers: mixed-folder upload, partial failures not aborting the batch,
duplicate detection (within and across jobs), meeting association by
folder, tenant isolation, cross-document retrieval, and security
(unauthorized access to another tenant's job, path traversal, oversized
file, invalid OAuth state is covered separately in test_auth_security.py).
"""
import asyncio
import io

import pytest
from docx import Document

from tests.conftest import dev_login

pytestmark = pytest.mark.asyncio

SAMPLE_VTT = (
    b"WEBVTT\n\n"
    b"00:00:00.000 --> 00:00:05.000\n<v Alice>We agreed to launch the pilot next week.</v>\n\n"
    b"00:00:05.500 --> 00:00:10.000\n<v Bob>I will prepare the demo.</v>\n"
)


def _sample_docx(heading: str, paragraph: str) -> bytes:
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _wait_for_job(client, headers, job_id: str, timeout_s: float = 10.0) -> dict:
    elapsed = 0.0
    while elapsed < timeout_s:
        resp = await client.get(f"/api/historical-imports/{job_id}", headers=headers)
        data = resp.json()
        if data["status"] not in ("queued", "processing"):
            return data
        await asyncio.sleep(0.1)
        elapsed += 0.1
    raise AssertionError(f"Import job {job_id} did not finish within {timeout_s}s")


async def test_mixed_folder_upload_indexes_vtt_and_docx_into_one_meeting(client):
    token = await dev_login(client, email="owner@acme.com", display_name="Owner", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [
        ("files", ("Project_A/Meeting_01.vtt", SAMPLE_VTT, "text/vtt")),
        ("files", ("Project_A/Architecture.docx", _sample_docx("Architecture", "API gateway notes."), "application/octet-stream")),
    ]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    assert create.status_code == 200
    job_id = create.json()["id"]

    job = await _wait_for_job(client, headers, job_id)
    assert job["status"] == "completed"
    assert job["successful_files"] == 2
    assert job["failed_files"] == 0

    meetings = (await client.get("/api/meetings", headers=headers)).json()
    assert len(meetings) == 1
    meeting = meetings[0]
    assert meeting["title"] == "Project A"
    assert meeting["is_historical"] is True
    assert meeting["document_count"] == 2
    assert meeting["transcript_available"] is True

    detail = (await client.get(f"/api/meetings/{meeting['id']}", headers=headers)).json()
    assert {d["file_type"] for d in detail["documents"]} == {"vtt", "docx"}


async def test_word_transcript_marks_meeting_ready(client):
    token = await dev_login(client, email="word@acme.com", display_name="Word", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    create = await client.post("/api/historical-imports", headers=headers, files=[
        ("files", ("Word/Video Transcript.docx", _sample_docx("Session", "Alice: Launch next week."), "application/octet-stream")),
    ])
    assert create.status_code == 200
    job = await _wait_for_job(client, headers, create.json()["id"])
    assert job["successful_files"] == 1
    meeting = (await client.get("/api/meetings", headers=headers)).json()[0]
    assert meeting["transcript_available"] is True
    assert meeting["status"] == "ready"
    detail = (await client.get(f"/api/meetings/{meeting['id']}", headers=headers)).json()
    assert detail["documents"][0]["document_type"] == "transcript"


async def test_one_unsupported_file_does_not_abort_the_batch(client):
    token = await dev_login(client, email="owner2@acme.com", display_name="Owner2", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [
        ("files", ("Project_B/Meeting_02.vtt", SAMPLE_VTT, "text/vtt")),
        ("files", ("Project_B/random.xyz", b"binary junk", "application/octet-stream")),
        ("files", ("Project_B/corrupt.docx", b"not a real docx", "application/octet-stream")),
    ]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    job_id = create.json()["id"]
    job = await _wait_for_job(client, headers, job_id)

    assert job["status"] == "completed_with_warnings"
    assert job["successful_files"] == 1
    assert job["skipped_files"] == 2
    assert job["failed_files"] == 0

    results = (await client.get(f"/api/historical-imports/{job_id}/results", headers=headers)).json()["results"]
    by_name = {r["filename"]: r for r in results}
    assert by_name["random.xyz"]["status"] == "skipped"
    assert by_name["random.xyz"]["recommended_action"]
    assert by_name["corrupt.docx"]["status"] == "skipped"
    assert by_name["Meeting_02.vtt"]["status"] == "success"


async def test_duplicate_file_is_detected_within_the_same_batch(client):
    token = await dev_login(client, email="owner3@acme.com", display_name="Owner3", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [
        ("files", ("Project_C/Meeting_03.vtt", SAMPLE_VTT, "text/vtt")),
        ("files", ("Project_C/Meeting_03_copy.vtt", SAMPLE_VTT, "text/vtt")),  # identical content, different name
    ]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    job_id = create.json()["id"]
    job = await _wait_for_job(client, headers, job_id)

    assert job["successful_files"] == 1
    assert job["skipped_files"] == 1
    results = (await client.get(f"/api/historical-imports/{job_id}/results", headers=headers)).json()["results"]
    statuses = sorted(r["status"] for r in results)
    assert statuses == ["duplicate", "success"]


async def test_reimporting_the_same_folder_is_fully_deduplicated(client):
    token = await dev_login(client, email="owner4@acme.com", display_name="Owner4", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    files = [("files", ("Project_D/Meeting_04.vtt", SAMPLE_VTT, "text/vtt"))]

    first = await client.post("/api/historical-imports", headers=headers, files=files)
    await _wait_for_job(client, headers, first.json()["id"])

    second = await client.post("/api/historical-imports", headers=headers, files=files)
    job2 = await _wait_for_job(client, headers, second.json()["id"])

    assert job2["successful_files"] == 0
    assert job2["skipped_files"] == 1
    meetings = (await client.get("/api/meetings", headers=headers)).json()
    assert len(meetings) == 1  # no duplicate meeting was created either


async def test_cross_document_retrieval_answers_from_transcript_and_docx(client):
    token = await dev_login(client, email="owner5@acme.com", display_name="Owner5", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}

    files = [
        ("files", ("Project_E/Meeting_05.vtt", SAMPLE_VTT, "text/vtt")),
        (
            "files",
            (
                "Project_E/Action_Items.docx",
                _sample_docx("Action Items", "Owner: Bob. Action: Prepare demo. Due Date: October 10."),
                "application/octet-stream",
            ),
        ),
    ]
    create = await client.post("/api/historical-imports", headers=headers, files=files)
    job = await _wait_for_job(client, headers, create.json()["id"])
    assert job["successful_files"] == 2

    meeting_id = (await client.get("/api/meetings", headers=headers)).json()[0]["id"]
    results = (
        await client.get(f"/api/meetings/{meeting_id}/search?q=demo", headers=headers)
    ).json()
    file_types = {r["file_type"] for r in results}
    assert "docx" in file_types  # the .docx chunk is retrievable alongside the transcript


async def test_chat_restores_imports_and_restricts_document_sources(client, monkeypatch):
    from types import SimpleNamespace
    from meeting_intel.agents import answer_agent
    from meeting_intel.llm.client import LLMResult
    from meeting_intel.retrieval.memory_search import get_memory_provider

    async def complete(**kwargs):
        prompt = str(kwargs)
        assert "SELECTED_TRANSCRIPT" in prompt
        assert "OTHER_DOCUMENT" not in prompt
        return LLMResult("The selected transcript discusses the pilot. [S1]", 1, "test")

    monkeypatch.setattr(answer_agent, "get_llm_provider", lambda: SimpleNamespace(complete=complete))
    token = await dev_login(client, email="scoped@acme.com", display_name="Scoped", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    created = await client.post("/api/historical-imports", headers=headers, files=[
        ("files", ("Scope/Transcript.docx", _sample_docx("Transcript", "SELECTED_TRANSCRIPT pilot launch."), "application/octet-stream")),
        ("files", ("Scope/Notes.docx", _sample_docx("Notes", "OTHER_DOCUMENT pilot budget."), "application/octet-stream")),
        ("files", ("Elsewhere/Other.docx", _sample_docx("Other", "OTHER_DOCUMENT private plan."), "application/octet-stream")),
    ])
    assert (await _wait_for_job(client, headers, created.json()["id"]))["successful_files"] == 3
    meetings = (await client.get("/api/meetings", headers=headers)).json()
    meeting = next(m for m in meetings if m["title"] == "Scope")
    detail = (await client.get(f"/api/meetings/{meeting['id']}", headers=headers)).json()
    document = next(d for d in detail["documents"] if d["source_file"] == "Transcript.docx")
    get_memory_provider().reset()  # simulate losing all process-local index data on reload
    response = await client.post("/api/chat", headers=headers, json={
        "meeting_id": meeting["id"], "document_id": document["id"],
        "message": "Across all meetings, what is this transcript about?",
    })
    assert response.status_code == 200, response.text
    assert response.json()["evidence_sufficient"] is True
    assert response.json()["cross_meeting"] is False
    assert {s["source_file"] for s in response.json()["sources"]} == {"Transcript.docx"}
    other = next(m for m in meetings if m["id"] != meeting["id"])
    denied = await client.post("/api/chat", headers=headers, json={
        "meeting_id": other["id"], "document_id": document["id"], "message": "pilot",
    })
    assert denied.status_code == 404
    other_token = await dev_login(client, email="outsider@example.com", display_name="Outsider", tenant_name="Other tenant")
    denied = await client.post("/api/chat", headers={"Authorization": f"Bearer {other_token}"}, json={
        "meeting_id": meeting["id"], "document_id": document["id"], "message": "pilot",
    })
    assert denied.status_code in (403, 404)


async def test_tenant_isolation_between_two_historical_imports(client):
    token_a = await dev_login(client, email="a@tenanta.com", display_name="A", tenant_name="TenantA")
    token_b = await dev_login(client, email="b@tenantb.com", display_name="B", tenant_name="TenantB")
    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    files = [("files", ("Confidential/Meeting.vtt", SAMPLE_VTT, "text/vtt"))]
    create = await client.post("/api/historical-imports", headers=headers_a, files=files)
    job_id = create.json()["id"]
    await _wait_for_job(client, headers_a, job_id)

    # Tenant B must not be able to see tenant A's import job or meeting.
    forbidden = await client.get(f"/api/historical-imports/{job_id}", headers=headers_b)
    assert forbidden.status_code == 404

    meetings_b = (await client.get("/api/meetings", headers=headers_b)).json()
    assert meetings_b == []


async def test_history_list_is_scoped_to_the_callers_tenant(client):
    token_a = await dev_login(client, email="a2@tenanta.com", display_name="A2", tenant_name="TenantA2")
    token_b = await dev_login(client, email="b2@tenantb.com", display_name="B2", tenant_name="TenantB2")
    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    files = [("files", ("F/m.vtt", SAMPLE_VTT, "text/vtt"))]
    create = await client.post("/api/historical-imports", headers=headers_a, files=files)
    await _wait_for_job(client, headers_a, create.json()["id"])

    history_a = (await client.get("/api/historical-imports", headers=headers_a)).json()
    history_b = (await client.get("/api/historical-imports", headers=headers_b)).json()
    assert len(history_a) == 1
    assert history_b == []


# --- Security ---


async def test_unauthenticated_import_is_rejected(client):
    files = [("files", ("F/m.vtt", SAMPLE_VTT, "text/vtt"))]
    resp = await client.post("/api/historical-imports", files=files)
    assert resp.status_code in (401, 403)


async def test_path_traversal_in_relative_path_is_rejected(client):
    token = await dev_login(client, email="sec1@acme.com", display_name="Sec1", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    files = [("files", ("../../etc/passwd", b"malicious", "text/plain"))]
    resp = await client.post("/api/historical-imports", headers=headers, files=files)
    assert resp.status_code == 400


async def test_oversized_file_is_rejected(client, monkeypatch):
    import meeting_intel.api.routers.historical_imports as router_module

    monkeypatch.setattr(router_module, "MAX_FILE_SIZE_BYTES", 10)
    token = await dev_login(client, email="sec2@acme.com", display_name="Sec2", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    files = [("files", ("F/big.vtt", b"x" * 1000, "text/vtt"))]
    resp = await client.post("/api/historical-imports", headers=headers, files=files)
    assert resp.status_code == 400


async def test_empty_upload_is_rejected(client):
    token = await dev_login(client, email="sec3@acme.com", display_name="Sec3", tenant_name="Acme")
    headers = {"Authorization": f"Bearer {token}"}
    resp = await client.post("/api/historical-imports", headers=headers, files=[])
    assert resp.status_code in (400, 422)
