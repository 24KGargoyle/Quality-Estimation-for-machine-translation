"""Unit tests for every `DocumentParser` implementation (§40 of the upgrade
spec): valid file, invalid/corrupted file, empty file, and — for legacy
formats — the "safely reported as unsupported" contract. No historical
import job or database is involved here; these test the parsers in
isolation, exactly as `ParserFactory.get_parser()` would hand them to the
import pipeline.
"""
import io

import pytest
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from pptx import Presentation

from meeting_intel.ingestion.documents import UnsupportedFileError
from meeting_intel.ingestion.parsers import ParserFactory, SUPPORTED_EXTENSIONS


def _kwargs(**overrides):
    base = dict(
        filename="file", relative_path="file", tenant_id="t1", meeting_id="m1",
        title="Title", document_id_prefix="doc",
    )
    base.update(overrides)
    return base


def test_all_expected_extensions_are_registered():
    for ext in ("vtt", "txt", "docx", "doc", "xlsx", "xls", "pdf", "pptx", "csv"):
        assert ext in SUPPORTED_EXTENSIONS
    assert ParserFactory.get_parser("unknownext") is None
    assert not ParserFactory.is_supported("zzz")


# --- VTT ---

class TestVTTParser:
    def test_valid_file(self):
        parser = ParserFactory.get_parser("vtt")
        content = b"WEBVTT\n\n00:00:00.000 --> 00:00:05.000\n<v Alice>Hello team.</v>\n"
        result = parser.parse(content=content, **_kwargs(filename="m.vtt", relative_path="m.vtt"))
        assert len(result.documents) == 1
        assert result.documents[0].speaker == "Alice"
        assert result.documents[0].document_type == "transcript"

    def test_empty_file(self):
        parser = ParserFactory.get_parser("vtt")
        result = parser.parse(content=b"", **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_invalid_content_no_crash(self):
        parser = ParserFactory.get_parser("vtt")
        result = parser.parse(content=b"not a vtt file at all", **_kwargs())
        assert result.documents == []


# --- TXT ---

class TestTextParser:
    def test_valid_file(self):
        parser = ParserFactory.get_parser("txt")
        result = parser.parse(content=b"Some plain notes about the project.", **_kwargs())
        assert len(result.documents) == 1
        assert result.documents[0].document_type == "supporting_document"

    def test_empty_file(self):
        parser = ParserFactory.get_parser("txt")
        result = parser.parse(content=b"   ", **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_large_file_is_chunked(self):
        parser = ParserFactory.get_parser("txt")
        big = "\n\n".join(f"Paragraph number {i} with some content in it." for i in range(200))
        result = parser.parse(content=big.encode(), **_kwargs())
        assert len(result.documents) > 1
        assert all(len(d.content) <= 900 for d in result.documents)


# --- DOCX / legacy DOC ---

class TestWordParser:
    def _sample_docx_bytes(self) -> bytes:
        doc = Document()
        doc.add_heading("Architecture", level=1)
        doc.add_paragraph("We discussed the API gateway design in detail today.")
        doc.add_heading("Action Items", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "Owner"
        table.rows[0].cells[1].text = "Action"
        table.rows[0].cells[2].text = "Due Date"
        row = table.add_row()
        row.cells[0].text = "Bob"
        row.cells[1].text = "Prepare demo"
        row.cells[2].text = "October 10"
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_valid_file_preserves_headings_and_tables(self):
        parser = ParserFactory.get_parser("docx")
        result = parser.parse(content=self._sample_docx_bytes(), **_kwargs(filename="Architecture.docx"))
        sections = {d.section for d in result.documents}
        assert "Architecture" in sections
        assert "Action Items" in sections
        table_doc = next(d for d in result.documents if d.metadata.get("is_table"))
        assert "Owner: Bob" in table_doc.content
        assert "Action: Prepare demo" in table_doc.content

    def test_empty_document(self):
        parser = ParserFactory.get_parser("docx")
        buf = io.BytesIO()
        Document().save(buf)
        result = parser.parse(content=buf.getvalue(), **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_corrupted_file_raises_unsupported(self):
        parser = ParserFactory.get_parser("docx")
        with pytest.raises(UnsupportedFileError):
            parser.parse(content=b"not a real docx", **_kwargs())

    def test_legacy_doc_reported_unsupported_not_crashed(self):
        parser = ParserFactory.get_parser("doc")
        with pytest.raises(UnsupportedFileError, match="Legacy .doc"):
            parser.parse(content=b"anything", **_kwargs(filename="old.doc"))


# --- XLSX / legacy XLS ---

class TestExcelParser:
    def _sample_xlsx_bytes(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Action Items"
        ws.append(["Owner", "Action", "Status", "Due Date"])
        for i in range(40):
            ws.append([f"Person{i}", f"Task {i}", "Open", "2026-01-01"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_valid_file_and_row_batching(self):
        parser = ParserFactory.get_parser("xlsx")
        result = parser.parse(content=self._sample_xlsx_bytes(), **_kwargs(filename="Project_Status.xlsx"))
        assert len(result.documents) == 2  # 40 rows / 25 per chunk -> 2 batches
        assert all(d.sheet_name == "Action Items" for d in result.documents)
        assert all(d.document_type == "spreadsheet" for d in result.documents)

    def test_empty_workbook(self):
        parser = ParserFactory.get_parser("xlsx")
        buf = io.BytesIO()
        Workbook().save(buf)
        result = parser.parse(content=buf.getvalue(), **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_corrupted_file_raises_unsupported(self):
        parser = ParserFactory.get_parser("xlsx")
        with pytest.raises(UnsupportedFileError):
            parser.parse(content=b"not a real xlsx", **_kwargs())

    def test_legacy_xls_reported_unsupported_not_crashed(self):
        parser = ParserFactory.get_parser("xls")
        with pytest.raises(UnsupportedFileError, match="Legacy .xls"):
            parser.parse(content=b"anything", **_kwargs(filename="old.xls"))


# --- PDF ---

class TestPDFParser:
    def _sample_pdf_bytes(self, text: str) -> bytes:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=14)
        pdf.cell(0, 10, text)
        return bytes(pdf.output())

    def test_valid_file(self):
        parser = ParserFactory.get_parser("pdf")
        result = parser.parse(
            content=self._sample_pdf_bytes("We agreed to launch the pilot next week."),
            **_kwargs(filename="Design.pdf"),
        )
        assert len(result.documents) == 1
        assert result.documents[0].page_number == 1
        assert "pilot" in result.documents[0].content

    def test_corrupted_file_raises_unsupported(self):
        parser = ParserFactory.get_parser("pdf")
        with pytest.raises(UnsupportedFileError):
            parser.parse(content=b"not a pdf", **_kwargs())

    def test_scanned_pdf_with_no_text_reports_ocr_required(self):
        parser = ParserFactory.get_parser("pdf")
        pdf = FPDF()
        pdf.add_page()  # a blank page has no extractable text — simulates a scanned image page
        blank = bytes(pdf.output())
        with pytest.raises(UnsupportedFileError, match="OCR"):
            parser.parse(content=blank, **_kwargs())


# --- PPTX ---

class TestPowerPointParser:
    def _sample_pptx_bytes(self) -> bytes:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "API Gateway Architecture"
        slide.placeholders[1].text = "Discussion of the gateway design"
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    def test_valid_file(self):
        parser = ParserFactory.get_parser("pptx")
        result = parser.parse(content=self._sample_pptx_bytes(), **_kwargs(filename="Architecture.pptx"))
        assert len(result.documents) == 1
        assert result.documents[0].slide_number == 1
        assert "API Gateway Architecture" in result.documents[0].content

    def test_empty_presentation(self):
        parser = ParserFactory.get_parser("pptx")
        buf = io.BytesIO()
        Presentation().save(buf)
        result = parser.parse(content=buf.getvalue(), **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_corrupted_file_raises_unsupported(self):
        parser = ParserFactory.get_parser("pptx")
        with pytest.raises(UnsupportedFileError):
            parser.parse(content=b"not a real pptx", **_kwargs())


# --- CSV ---

class TestCSVParser:
    def test_valid_file(self):
        parser = ParserFactory.get_parser("csv")
        content = b"Owner,Action,Status\nBob,Prepare demo,In Progress\nAlice,Review demo,Pending\n"
        result = parser.parse(content=content, **_kwargs(filename="Action_Items.csv"))
        assert len(result.documents) == 1
        assert "Bob" in result.documents[0].content

    def test_large_file_is_batched(self):
        parser = ParserFactory.get_parser("csv")
        rows = "\n".join(f"p{i},task {i},open" for i in range(60))
        content = f"Owner,Action,Status\n{rows}\n".encode()
        result = parser.parse(content=content, **_kwargs())
        assert len(result.documents) == 3  # 60 rows / 25 per chunk -> 3 batches

    def test_header_only_no_data_rows(self):
        parser = ParserFactory.get_parser("csv")
        result = parser.parse(content=b"Owner,Action,Status\n", **_kwargs())
        assert result.documents == []
        assert result.warnings

    def test_empty_file(self):
        parser = ParserFactory.get_parser("csv")
        result = parser.parse(content=b"", **_kwargs())
        assert result.documents == []
        assert result.warnings
